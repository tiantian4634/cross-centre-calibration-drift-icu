# -*- coding: utf-8 -*-
"""生成 STROBE 官方格式清单（投稿版）—— Markdown + 可编辑 DOCX。

格式依据：STROBE Statement (von Elm et al., 2007) 官方 checklist 表头
          Item No | Recommendation | Page No. | Relevant text from manuscript
本文件把「Page No.」列扩为「Section / Page No.」，因为页数取决于作者最终排版，
先给出稳定可定位的章节锚点，页码由作者填写。
"""
import io, sys, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

B = r"E:\文件杂\褚俊智老师\论文选题\选题T1_虚弱营养列线图\09_路B_跨院校准"
OUT_MD = os.path.join(B, "07_投稿附件", "STROBE_checklist_投稿版.md")
OUT_DOCX = os.path.join(B, "07_投稿附件", "STROBE_checklist_投稿版.docx")

TITLE = ("Discrimination transports, calibration does not: calibration drift in severity "
         "scoring across 187 intensive-care units")

# ── 条目数据：(No, Recommendation, Section anchor, Relevant text) ─────────────
S = [
("1a", "Indicate the study's design with a commonly used term in the title or the abstract.",
 "Abstract, sentence 1",
 "\u201cIn a secondary analysis of eICU-CRD (30,796 patients aged \u226565 years from 187 hospitals) we examined whether such scores discriminate and calibrate alike across centres.\u201d"),

("1b", "Provide in the abstract an informative and balanced summary of what was done and what was found.",
 "Abstract (200/200 words, unstructured)",
 "Reports the main estimate (\u201cpooled AUC 0.732 \u2026 between-centre I\u00b2 was 75.4% and the observed-to-expected mortality ratio ranged from 0.69 to 1.54, a 2.2-fold spread\u201d), the recalibration penalty (\u201cwith 25 local patients it performed worse than no recalibration (success rate 15.5% versus 21.4%)\u201d), and the shrinkage result (\u201coutperforming naive recalibration at 25 patients (24.3% versus 15.5%; two-sided P = 3.7 \u00d7 10\u207b\u00b9\u00b9\u00b2)\u201d). Two negative findings are stated."),

("2", "Explain the scientific background and rationale for the investigation being reported.",
 "Introduction \u00b61\u20133",
 "\u201cCalibration \u2014 whether the predicted probability of death matches the observed probability \u2014 is the property that determines what a risk estimate actually means at the bedside \u2026\u201d"),

("3", "State specific objectives, including any prespecified hypotheses.",
 "Introduction, final \u00b6",
 "\u201cWe quantified between-centre variation in calibration for three established severity scores, tested how much of it measured case-mix could explain, decomposed it into components attributable to centre size and to calendar year, translated it into consequences at four decision thresholds, and evaluated local recalibration as a function of the number of locally available patients \u2014 together with an empirical-Bayes shrinkage alternative.\u201d"),

("4", "Present key elements of study design early in the paper.",
 "Methods \u2192 Study design and setting",
 "\u201cWe performed a retrospective secondary analysis of two publicly available, de-identified, multicentre critical-care databases: the eICU Collaborative Research Database (eICU-CRD v2.0) and MIMIC-IV v3.1.\u201d"),

("5", "Describe the setting, locations, and relevant dates, including periods of recruitment, exposure, follow-up, and data collection.",
 "Methods \u2192 Data sources",
 "\u201ceICU-CRD v2.0 contains data from 208 hospitals in the United States for the period 2014\u20132015, comprising 200,859 intensive-care unit (ICU) stays.\u201d \u2026 \u201cMIMIC-IV v3.1 contains data from Beth Israel Deaconess Medical Center, a single academic medical centre, comprising 94,458 ICU stays.\u201d"),

("6a", "Give the eligibility criteria, and the sources and methods of selection of participants. Describe methods of follow-up.",
 "Methods \u2192 Cohort construction; Fig. 1",
 "\u201cOf 200,859 unit stays across 208 hospitals, we retained the first ICU stay per patient (139,367), restricted to patients aged \u2265 65 years (71,734) with an ICU length of stay \u2265 24 h (50,138) and admitted to a medical ICU or medical\u2013surgical ICU (31,307).\u201d Centre-level eligibility: \u201ca hospital had to contribute \u2265 100 patients and \u2265 10 in-hospital deaths \u2026 103 of 187 hospitals met it.\u201d"),

("6b", "For matched studies, give matching criteria and the number of exposed and unexposed.",
 "N/A",
 "Not a matched design. No matching was performed; all eligible patients were included (see 6a)."),

("7", "Clearly define all outcomes, exposures, predictors, potential confounders, and effect modifiers. Give diagnostic criteria, if applicable.",
 "Methods \u2192 Severity scores; \u2192 Centre-level calibration measures",
 "\u201cIn-hospital death is a discharge status recorded in the source databases rather than a derived outcome, and no case-finding algorithm was applied to it.\u201d The two centre-level quantities are defined in full: \u201cwe solved for a_h such that the mean of expit(logit(p) + a_h) equals \u0177, by bisection over [\u221212, 12] to a tolerance of 10\u207b\u2074\u201d; and the calibration slope b_h."),

("8", "For each variable of interest, give sources of data and details of methods of assessment (measurement). Describe comparability of assessment methods if there is more than one group.",
 "Methods \u2192 Data sources; \u2192 Severity scores",
 "\u201cThree established severity scores were used exactly as supplied by the source extraction protocols \u2026 This study did not recompute the scores, and their derivation is documented in the source protocols.\u201d"),

("9", "Describe any efforts to address potential sources of bias.",
 "Methods \u2192 Potential sources of bias",
 "\u201cFour sources were considered.\u201d Selection (24-h stay and medical-ICU restriction; 24-h deaths excluded), information (scores and outcome not re-adjudicated), leakage (\u201ceach hospital's shrinkage parameters were estimated from the remaining hospitals only\u201d), and confounding of the between-centre comparison (\u201cNo causal comparison between hospitals is made or intended\u201d)."),

("10", "Explain how the study size was arrived at.",
 "Methods \u2192 Study size and missing data",
 "\u201cNo sample size calculation was performed. This is a secondary analysis of two existing databases, and the analysis used all patients meeting the eligibility criteria described above.\u201d"),

("11", "Explain how quantitative variables were handled in the analyses. If applicable, describe which groupings were chosen and why.",
 "Methods \u2192 Quantitative variables and cut-points",
 "\u201cAll three scores, age and comorbidity burden were analysed as continuous variables and standardised before fitting; they were not categorised. Cut-points used in this study were fixed in advance and are not data-driven \u2026\u201d (all thresholds listed and traced to the pre-specified plans)."),

("12a", "Describe all statistical methods, including those used to control for confounding.",
 "Methods \u2192 Pooled models through \u2192 Statistical reporting",
 "\u201cFor each score we fitted a pooled logistic regression across the full cohort, with the score standardised (z-score) before fitting.\u201d Plus: random-effects meta-analysis with DerSimonian\u2013Laird \u03c4\u00b2, leave-one-centre-out cross-fitting, paired two-sided Wilcoxon signed-rank tests, and cluster bootstrap for pooled AUC confidence intervals."),

("12b", "Describe any methods used to examine subgroups and interactions.",
 "Methods \u2192 Centre size and clinical impact; \u2192 Case-mix ladder",
 "\u201cHospitals were split into tertiles by sample size, and I\u00b2 and \u03c4\u00b2 were recomputed within each tertile.\u201d Three nested case-mix models (M1\u2013M3) were also compared. No interaction test was performed; subgroup analyses are descriptive."),

("12c", "Explain how missing data were addressed.",
 "Methods \u2192 Study size and missing data; Supplementary Table S1",
 "\u201cIn the eICU-CRD cohort the three severity scores, age, sex, invasive ventilation and in-hospital death were complete in every patient (0% missing), so no imputation was performed. Charlson comorbidity index was missing in 3,560 of 30,796 patients (11.6%) in eICU-CRD \u2026 and in 3 of 8,579 patients (0.04%) in MIMIC-IV.\u201d Comorbidity enters only the descriptive M3, which uses complete cases."),

("12d", "If applicable, explain how loss to follow-up was addressed.",
 "N/A",
 "The outcome is in-hospital death, taken from the discharge status recorded in the source database. There is no post-discharge follow-up and therefore no loss to follow-up."),

("12e", "Describe any sensitivity analyses.",
 "Methods \u2192 Between-centre heterogeneity; \u2192 Temporal stability; \u2192 Analysis framework",
 "\u201cBecause DerSimonian\u2013Laird predates current guidance, \u03c4\u00b2 was additionally re-estimated under three further frameworks as a pre-specified analysis (analysis K): Paule\u2013Mandel, restricted maximum likelihood (REML), and a one-stage random-intercept multilevel logistic model \u2026\u201d The case-mix ladder (M1\u2013M3) and a raw-scale (unstandardised) refit serve as further sensitivity analyses."),

("13a", "Report the numbers of individuals at each stage of study \u2014 e.g. numbers potentially eligible, examined for eligibility, confirmed eligible, included in the study, completing follow-up, and analysed.",
 "Results \u2192 Cohort and study setting; Fig. 1",
 "\u201cThe eICU-CRD cohort comprised 30,796 patients from 187 hospitals, of whom 3,901 (12.7%) died in hospital; 103 hospitals met the pre-specified centre-level eligibility criterion of \u2265 100 patients and \u2265 10 in-hospital deaths and entered the centre-level analyses \u2026 The MIMIC-IV cohort comprised 8,579 patients from a single academic medical centre, of whom 1,654 (19.3%) died (Table 1; Supplementary Table S1).\u201d All intermediate counts are in Fig. 1 and Methods."),

("13b", "Give reasons for non-participation at each stage.",
 "Fig. 1; Methods \u2192 Cohort construction",
 "\u201cWe then excluded deaths within 24 h of ICU admission (n = 7), stays with an unknown hospital outcome status (n = 363), and stays in which albumin, height and weight were all missing (n = 141).\u201d Reasons are given for every exclusion step in both cohorts."),

("13c", "Consider use of a flow diagram.",
 "Fig. 1",
 "\u201cThe step-wise inclusion and exclusion criteria are shown in Figure 1, and the counts reported there are taken from the extraction stage's own count reports \u2026 rather than re-derived, so that Figure 1 reflects the actual pipeline rather than a reconstruction.\u201d"),

("14a", "Give characteristics of study participants (eg demographic, clinical, social) and information on exposures and potential confounders.",
 "Table 1; Supplementary Table S1",
 "Table 1 reports, by score and cohort, the number of patients and hospitals, age, sex, invasive ventilation, Charlson index, observed mortality and the pooled discrimination and calibration estimates. Supplementary Table S1 gives the full cohort characteristics with a missing-data row for every variable."),

("14b", "Indicate the number of participants with missing data for each variable of interest.",
 "Supplementary Table S1; Methods \u2192 Study size and missing data",
 "\u201cCharlson comorbidity index was missing in 3,560 of 30,796 patients (11.6%) in eICU-CRD (2,807 of 26,436, 10.6%, within the 103-centre analysis subset) and in 3 of 8,579 patients (0.04%) in MIMIC-IV.\u201d All other analysis variables were 0% missing."),

("14c", "Summarise follow-up time (eg, average and total amount).",
 "Methods \u2192 Cohort construction; \u2192 Study size and missing data",
 "The observation window is the index hospital admission; \u201cIn-hospital death is a discharge status recorded in the source databases\u201d. ICU length of stay \u2265 24 h was an eligibility criterion, so no patient contributed less than one day of observation."),

("15", "Report numbers of outcome events or summary measures over time.",
 "Results \u2192 Cohort and study setting; Table 1; Supplementary Table S1",
 "eICU-CRD: \u201c3,901 (12.7%) died in hospital\u201d of 30,796. MIMIC-IV: \u201c1,654 (19.3%) died\u201d of 8,579. Per-hospital event counts are the input to every centre-level estimate and are recoverable from the archived code."),

("16a", "Give unadjusted estimates and, if applicable, confounder-adjusted estimates and their precision (eg, 95% confidence interval). Make clear which confounders were adjusted for and why they were included.",
 "Results \u2192 Calibration varies widely across centres; Table 1; Fig. 3",
 "\u201cFor SOFA, the pooled model discriminated moderately (AUC 0.732)\u201d (95% CI 0.720\u20130.745 by cluster bootstrap, 1,000 replicates). Covariates enter only the case-mix ladder, where the reason is stated: \u201cTo test whether measured case-mix explains between-centre calibration drift\u201d. Precision is reported for every estimate."),

("16b", "Report category boundaries when continuous variables were categorised.",
 "Methods \u2192 Quantitative variables and cut-points",
 "\u201cthe centre-size tertile boundaries (hospitals split by sample size into three groups of 36, 33 and 34 hospitals, with ranges 100\u2013181, 184\u2013280 and 289\u2013987 patients respectively)\u201d. The four decision thresholds (0.10, 0.20, 0.30, 0.40) are also given. Scores and age were not categorised."),

("16c", "If relevant, consider translating estimates of relative risk into absolute risk for a meaningful time period.",
 "N/A",
 "No relative-risk estimates are reported. The study reports absolute risks directly \u2014 observed mortality, predicted mortality, the observed-to-expected ratio, and mortality among patients flagged above each decision threshold."),

("17", "Report other analyses done \u2014 eg analyses of subgroups and interactions, and sensitivity analyses.",
 "Results \u2192 Measured case-mix; \u2192 Drift strengthens; \u2192 Calibration offset is not a stable attribute; \u2192 The between-centre variance depends on the estimation framework; Supplementary Fig. S1",
 "Reports the case-mix ladder with its two-way decomposition (\u201c1.9 points were therefore attributable to the change in analysis set and 3.4 points to the addition of the covariates themselves\u201d), the size tertiles, the split-half temporal analysis, and \u03c4\u00b2 under four frameworks."),

("18", "Summarise key results with reference to study objectives.",
 "Discussion \u00b61",
 "\u201cIn 103 hospitals using the same scores and the same pipeline, a SOFA-based model discriminated consistently (per-centre AUC 0.649\u20130.812) while its calibration varied by more than twofold (observed-to-expected mortality 0.69\u20131.54; between-centre I\u00b2 = 75.4%).\u201d"),

("19", "Discuss limitations of the study, taking into account sources of potential bias or imprecision. Discuss both the direction and magnitude of any potential bias.",
 "Discussion \u2192 Limitations",
 "\u201cFirst, both databases were collected for purposes other than this study \u2026 Misclassification in the underlying records \u2026 cannot be detected here, and its direction and magnitude could not be quantified.\u201d Six limitations in total, each with its direction of effect where the direction is knowable."),

("20", "Give a cautious overall interpretation of results considering objectives, limitations, multiplicity of analyses, results from similar studies, and other relevant evidence.",
 "Discussion \u00b61\u20136; Methods \u2192 Analysis framework and inferential status",
 "Findings are set against prior calibration-drift and model-updating work, with the qualifying rather than contradicting relation made explicit. Multiplicity is addressed directly: \u201cno formal adjustment for multiple comparisons was applied\u201d, with analyses assigned to pre-specified confirmatory or descriptive tiers and two unmet criteria reported as negative findings."),

("21", "Discuss the generalisability (external validity) of the study results.",
 "Discussion \u2192 Limitations, final paragraph",
 "\u201cthe empirical-Bayes parameters were estimated in a single database, and although they transferred to a second, the extent to which they generalise to other health systems and other score families remains open.\u201d The scope of inference is further bounded by the 2014\u20132015 eICU-CRD window."),

("22", "Give the source of funding and the role of the funders for the present study and, if applicable, for the original study on which the present article is based.",
 "Additional Information \u2192 Funding",
 "\u201cThis research received no external funding.\u201d No funder had any role in design, analysis, interpretation, or the decision to submit."),
]

# ── RECORD 扩展（13 条）────────────────────────────────────
R = [
("R1.1", "The type of data and the name of the database(s) used should be reported in the title or abstract.",
 "Abstract", "\u201cIn a secondary analysis of eICU-CRD \u2026\u201d and \u201c\u2026 transferring to MIMIC-IV without refitting.\u201d Both databases are named in the abstract."),
("R1.2", "The geographic region and timeframe within which the study took place should be reported in the title or abstract.",
 "Abstract (not reported); Methods \u2192 Data sources",
 "Not in the abstract, which is at its 200-word limit. Region and period are reported in Methods: United States, 2014\u20132015 (eICU-CRD); single US academic centre (MIMIC-IV). Author decision, recorded here."),
("R1.3", "If the study used a linkage between databases, this should be described in the title or abstract.",
 "N/A", "No record linkage between databases. The two cohorts are analysed independently."),
("R6.1", "The methods of selection of the study population using codes or algorithms should be reported.",
 "Methods \u2192 Cohort construction; Fig. 1",
 "Every inclusion and exclusion step is reported with its count and its criterion. No code list was required, because eligibility rests on age, ICU type, length of stay and recorded discharge status rather than on diagnostic coding."),
("R6.2", "The validation of the codes or algorithms used to select the population should be referenced.",
 "Methods \u2192 Validity of the variable definitions",
 "\u201cSOFA, OASIS and SAPS-II are established, previously published and externally validated scores, used here exactly as supplied by the source extraction protocols; their components and weights are those of the original publications and are not modified in this study.\u201d"),
("R6.3", "If the study involved linkage of databases, a flow diagram of the linkage should be provided.",
 "N/A", "No linkage; consequently no linkage flow diagram. Fig. 1 shows participant flow for each cohort separately."),
("R7.1", "A complete list of codes and algorithms used to classify exposures, outcomes, confounders, and effect modifiers should be provided.",
 "Methods \u2192 Validity of the variable definitions; \u2192 Centre-level calibration measures",
 "\u201cNo code list or phenotype algorithm was developed for this study; where a variable required a derived definition, that definition is given above in full (the calibration offset and its standard error, and the calibration slope).\u201d Both derived definitions are given as closed-form expressions."),
("R12.1", "Describe the extent to which the investigators had access to the database population.",
 "Methods \u2192 Data sources", "\u201cAll database access was read-only; no script written for this study modified a database.\u201d Access was to the complete credentialed databases, not to a restricted extract."),
("R12.2", "Describe the methods used to clean the data.",
 "Methods \u2192 Cohort construction; \u2192 Study size and missing data",
 "Cleaning is confined to the reported eligibility filters; no value was corrected, imputed or winsorised. Missingness is quantified for every variable rather than handled by a cleaning rule."),
("R12.3", "Describe the methods used to link databases, and the quality of the linkage.",
 "N/A", "No linkage was performed, so linkage quality does not apply."),
("R13.1", "Describe the selection of participants, including data-quality and completeness filters.",
 "Methods \u2192 Cohort construction; Fig. 1",
 "Includes the data-availability filters: \u201cstays with an unknown hospital outcome status (n = 363)\u201d and \u201cstays in which albumin, height and weight were all missing (n = 141)\u201d, each with its count and its reason."),
("R19.1", "Discuss the implications of using data that were not created to answer the research question.",
 "Discussion \u2192 Limitations, first limitation",
 "\u201cMisclassification in the underlying records \u2014 in the severity scores, in comorbidity coding and in the recorded discharge status \u2014 cannot be detected here, and its direction and magnitude could not be quantified. And because eICU-CRD covers only 2014 and 2015, we could not assess how eligibility for the cohort, or the coding practices behind it, changed over time.\u201d"),
("R22.1", "Authors should provide information on how to access the study protocol, raw data, and programming code.",
 "Data Availability; Methods \u2192 Software, seed and reproducibility",
 "Analysis plans and their timestamps are reproduced in the Supplementary Note; code and derived results are archived with a DOI; raw data are available from PhysioNet under the original data use agreements."),
]

# ── Markdown ────────────────────────────────────────────────────────────────
md = []
md.append("<!-- Source file for STROBE_checklist_投稿版.docx. Edit here, then re-run 生成_STROBE清单.py -->")
md.append("")
md.append("# STROBE Statement \u2014 checklist of items that should be included in reports of cohort studies")
md.append("")
md.append("| Field | Value |")
md.append("|---|---|")
md.append("| **Manuscript title** | %s |" % TITLE)
md.append("| **Study design** | Retrospective secondary analysis of two routinely collected critical-care databases (cohort design) |")
md.append("| **Guideline** | STROBE 2007, cohort-study checklist; RECORD 2015 extension (Part B) |")
md.append("| **Journal** | Scientific Reports |")
md.append("| **Item count** | 22 STROBE items (34 sub-items) + 13 RECORD items |")
md.append("")
md.append("> **How to complete this form.** The \u201cSection\u201d column gives a stable anchor \u2014 a heading name that does not "
          "change with pagination. Fill the page number in the \u201cPage No.\u201d column from the final typeset manuscript, "
          "then export to PDF for upload. The \u201cRelevant text\u201d column quotes the manuscript verbatim so that each "
          "claim can be located by search as well as by page.")
md.append("")
md.append("## Part A \u2014 STROBE (cohort studies)")
md.append("")
md.append("| Item No | Recommendation | Section | Page No. | Relevant text from manuscript |")
md.append("|---|---|---|---|---|")
for no, rec, sec, txt in S:
    md.append("| **%s** | %s | %s | | %s |" % (no, rec, sec, txt))
md.append("")
md.append("## Part B \u2014 RECORD extension (routinely collected health data)")
md.append("")
md.append("| Item No | Recommendation | Section | Page No. | Relevant text from manuscript |")
md.append("|---|---|---|---|---|")
for no, rec, sec, txt in R:
    md.append("| **%s** | %s | %s | | %s |" % (no, rec, sec, txt))
md.append("")
md.append("## Notes on items recorded as N/A")
md.append("")
md.append("| Item | Reason for N/A |")
md.append("|---|---|")
md.append("| 6b | Not a matched design; no matching was performed. |")
md.append("| 12d | Outcome is in-hospital death recorded at discharge; no post-discharge follow-up, hence no loss to follow-up. |")
md.append("| 16c | No relative-risk estimates; absolute risks and observed-to-expected ratios are reported directly. |")
md.append("| R1.3, R6.3, R12.3 | The two databases were analysed independently; no record linkage was performed. |")
md.append("")
open(OUT_MD, "w", encoding="utf-8").write("\n".join(md) + "\n")
print("  \u2714 %s" % os.path.basename(OUT_MD))

# ── DOCX ────────────────────────────────────────────────────────────────────
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.6)
sec.top_margin = sec.bottom_margin = Cm(1.6)

st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(9)
st._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("STROBE Statement\u2014checklist of items that should be included in reports of cohort studies")
r.bold = True; r.font.size = Pt(13)
p.paragraph_format.space_after = Pt(2)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(TITLE)
r.italic = True; r.font.size = Pt(9.5)
p.paragraph_format.space_after = Pt(8)

# 信息表
info = [("Guideline", "STROBE 2007 (cohort studies); RECORD 2015 extension in Part B"),
        ("Study design", "Retrospective secondary analysis of two routinely collected critical-care databases"),
        ("Journal", "Scientific Reports"),
        ("Items", "22 STROBE items (34 sub-items) + 13 RECORD items"),
        ("Section anchor", "Heading names are stable across pagination; fill page numbers from the final typeset file")]
tb = doc.add_table(rows=0, cols=2)
tb.style = "Table Grid"
tb.alignment = WD_TABLE_ALIGNMENT.CENTER
for k, v in info:
    row = tb.add_row().cells
    row[0].width = Cm(3.4); row[1].width = Cm(14.4)
    rr = row[0].paragraphs[0].add_run(k); rr.bold = True; rr.font.size = Pt(8.5)
    rr2 = row[1].paragraphs[0].add_run(v); rr2.font.size = Pt(8.5)
for row in tb.rows:
    row.cells[0].width = Cm(3.4); row.cells[1].width = Cm(14.4)

doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_part(title, items):
    h = doc.add_paragraph()
    hr = h.add_run(title)
    hr.bold = True; hr.font.size = Pt(10.5)
    h.paragraph_format.space_before = Pt(6)
    h.paragraph_format.space_after = Pt(3)

    t = doc.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Cm(1.3), Cm(5.7), Cm(3.6), Cm(1.5), Cm(5.7)]
    hdr = ["Item No", "Recommendation", "Section", "Page No.", "Relevant text from manuscript"]
    for i, hname in enumerate(hdr):
        c = t.rows[0].cells[i]
        c.width = widths[i]
        pp = c.paragraphs[0]
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = pp.add_run(hname); rr.bold = True; rr.font.size = Pt(8.5)

    # 表头跨页重复
    trPr = t.rows[0]._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true"); trPr.append(el)

    for no, rec, secname, txt in items:
        cells = t.add_row().cells
        for i in range(5):
            cells[i].width = widths[i]
        # Item No
        pp = cells[0].paragraphs[0]
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = pp.add_run(no); rr.bold = True; rr.font.size = Pt(8)
        # Recommendation
        rr = cells[1].paragraphs[0].add_run(rec); rr.font.size = Pt(8)
        # Section
        rr = cells[2].paragraphs[0].add_run(secname); rr.font.size = Pt(8)
        # Page (blank, lightly marked)
        pp = cells[3].paragraphs[0]
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = pp.add_run(""); rr.font.size = Pt(8)
        # Relevant text
        rr = cells[4].paragraphs[0].add_run(txt); rr.font.size = Pt(8)

    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = w


add_part("Part A \u2014 STROBE (cohort studies)", S)
doc.add_page_break()
add_part("Part B \u2014 RECORD extension (routinely collected health data)", R)

# N/A 说明
h = doc.add_paragraph()
hr = h.add_run("Notes on items recorded as N/A")
hr.bold = True; hr.font.size = Pt(10)
h.paragraph_format.space_before = Pt(10)

na = [("6b", "Not a matched design; no matching was performed."),
      ("12d", "Outcome is in-hospital death recorded at discharge; no post-discharge follow-up, hence no loss to follow-up."),
      ("16c", "No relative-risk estimates; absolute risks and observed-to-expected ratios are reported directly."),
      ("R1.3 / R6.3 / R12.3", "The two databases were analysed independently; no record linkage was performed.")]
nt = doc.add_table(rows=1, cols=2)
nt.style = "Table Grid"
for i, hname in enumerate(["Item", "Reason"]):
    c = nt.rows[0].cells[i]
    c.width = Cm(3.6) if i == 0 else Cm(14.2)
    rr = c.paragraphs[0].add_run(hname); rr.bold = True; rr.font.size = Pt(8.5)
for k, v in na:
    cells = nt.add_row().cells
    cells[0].width = Cm(3.6); cells[1].width = Cm(14.2)
    rr = cells[0].paragraphs[0].add_run(k); rr.bold = True; rr.font.size = Pt(8)
    rr = cells[1].paragraphs[0].add_run(v); rr.font.size = Pt(8)

doc.save(OUT_DOCX)
print("  \u2714 %s  (%.1f KB)" % (os.path.basename(OUT_DOCX), os.path.getsize(OUT_DOCX) / 1024))
print("\n  STROBE 子条目 %d ｜ RECORD 条目 %d ｜ 合计 %d" % (len(S), len(R), len(S) + len(R)))
