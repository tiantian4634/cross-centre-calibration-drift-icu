# -*- coding: utf-8 -*-
"""投稿包 DOCX 质量检查：Markdown 是否残留、表格是否为真 Word 表、占位符是否可见。"""
import io, sys, os, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from docx import Document

B = r"E:\文件杂\褚俊智老师\论文选题\选题T1_虚弱营养列线图\09_路B_跨院校准"
PKG = os.path.join(B, "00_投稿包")
docx_path = os.path.join(PKG, "1_稿件_Manuscript", "Manuscript.docx")

d = Document(docx_path)
paras = [p.text for p in d.paragraphs]
allp = "\n".join(paras)
tbl_txt = "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
full = allp + "\n" + tbl_txt

print("  ── 结构 ──")
print("    段落 %d ｜ 表格 %d 个" % (len(d.paragraphs), len(d.tables)))
for k, t in enumerate(d.tables, 1):
    print("      表 %d：%d 行 × %d 列" % (k, len(t.rows), len(t.columns)))

print("\n  ── 标题层级 ──")
hs = {}
for p in d.paragraphs:
    if p.style.name.startswith("Heading") and p.text.strip():
        hs.setdefault(p.style.name, []).append(p.text.strip())
for k in sorted(hs):
    print("    %-10s %d 个" % (k, len(hs[k])))
for k in sorted(hs):
    print("      %s 示例：%s" % (k, hs[k][0][:60]))

print("\n  ── Markdown 残留检测（应全为 0）──")
checks = [
    ("独立 | 管线（表格未解析）", len(re.findall(r"(?m)^\s*\|", full))),
    ("** 粗体标记", full.count("**")),
    ("反引号 `", full.count("`")),
    ("行首 # 标题标记", len(re.findall(r"(?m)^#{1,6}\s", full))),
    ("行首 > 引用标记", len(re.findall(r"(?m)^>\s", full))),
    ("--- 分隔线", len(re.findall(r"(?m)^---\s*$", full))),
    ("⟨ 未高亮占位符", full.count("⟨")),
]
for name, v in checks:
    print("    %-30s %d  %s" % (name, v, "✔" if v == 0 or name.startswith("⟨") else "✘"))

print("\n  ── 关键内容是否到位 ──")
keys = {
    "标题": "Discrimination transports, calibration does not",
    "摘要首句": "Severity scores are often assumed to transport",
    "Data Availability 段落": "zenodo.23479141",
    "Funding": "received no external funding",
    "Author contributions": "CRediT",
    "Competing interests": "Competing interests",
    "Figure legends": "Figure 1",
    "Table 1": "Table 1",
    "References": "DerSimonian",
}
for k, v in keys.items():
    print("    %-24s %s" % (k, "✔" if v in full else "✘ 缺失"))

print("\n  ── 顺序检查（Sci Rep 要求）──")
def pos(s):
    return allp.find(s)
order = [("## Data Availability", "Data Availability"), ("References", "References"),
         ("Author contributions", "Author contributions"),
         ("Additional Information", "Additional Information"),
         ("Figure legends", "Figure legends")]
idx = [(n, pos(t)) for n, t in order]
for n, i in idx:
    print("    %-26s 位置 %d" % (n, i))
ok = all(idx[k][1] < idx[k + 1][1] for k in range(len(idx) - 1))
print("    顺序正确：%s" % ("✔" if ok and idx[0][1] > 0 else "✘"))

# 文件清单
print("\n  ── 包内文件 ──")
for root, dirs, files in os.walk(PKG):
    if "_manuscript.html" in files or "_conv.ps1" in files:
        pass
    for f in sorted(files):
        p = os.path.join(root, f)
        rel = os.path.relpath(p, PKG)
        print("    %-62s %8.0f KB" % (rel, os.path.getsize(p) / 1024))
